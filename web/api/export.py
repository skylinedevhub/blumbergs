"""POST /api/export — Execute query and return .xlsx file."""

import json
import os
import re
import io
import decimal
from http.server import BaseHTTPRequestHandler
import psycopg2
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font


DATABASE_URL = os.environ.get("NEON_DATABASE_URL", "")


def prep(sql):
    sql = re.sub(
        r"(--[^\n]*\n\s*)?CREATE\s+OR\s+REPLACE\s+MACRO\s+money.*?;\s*", "", sql, flags=re.S | re.I
    )
    return sql.strip()


def only_select(sql):
    s = re.sub(r"--[^\n]*", "", sql).strip().upper()
    if not (s.startswith("SELECT") or s.startswith("WITH")):
        return False
    for kw in (
        "INSERT", "UPDATE", "DELETE", "DROP", "ALTER",
        "CREATE", "TRUNCATE", "COPY", "ATTACH", "DETACH",
    ):
        if re.search(r"\b" + kw + r"\b", s):
            return False
    return True


class handler(BaseHTTPRequestHandler):
    def do_POST(self):
        try:
            body = self.rfile.read(int(self.headers.get("Content-Length", 0))).decode()
            data = json.loads(body)
            sql = prep(data.get("sql", ""))
            sql = re.sub(r"\bLIMIT\s+\d+\s*;?\s*$", ";", sql, flags=re.I)

            if not only_select(sql):
                return self._json({"error": "Only SELECT queries allowed"}, 400)

            con = psycopg2.connect(DATABASE_URL)
            cur = con.cursor()
            cur.execute(sql)
            cols = [d[0] for d in cur.description]
            rows = cur.fetchall()
            cur.close()
            con.close()

            wb = Workbook()
            ws = wb.active
            ws.title = "T3010 Export"
            hfont = Font(bold=True)
            for ci, col in enumerate(cols, 1):
                cell = ws.cell(row=1, column=ci, value=col)
                cell.font = hfont

            for ri, row in enumerate(rows, 2):
                for ci, val in enumerate(row, 1):
                    if val is None:
                        continue
                    if isinstance(val, decimal.Decimal):
                        val = float(val)
                    ws.cell(row=ri, column=ci, value=val)

            for ci in range(1, len(cols) + 1):
                sample = [len(str(cols[ci - 1]))]
                for ri in range(min(50, len(rows))):
                    sample.append(len(str(rows[ri][ci - 1] or "")))
                ws.column_dimensions[get_column_letter(ci)].width = min(max(sample) + 2, 45)

            buf = io.BytesIO()
            wb.save(buf)
            raw = buf.getvalue()

            self.send_response(200)
            self.send_header(
                "Content-Type",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            self.send_header(
                "Content-Disposition", 'attachment; filename="t3010_export.xlsx"'
            )
            self.send_header("Content-Length", str(len(raw)))
            self.end_headers()
            self.wfile.write(raw)
        except Exception as e:
            self._json({"error": str(e)}, 500)

    def _json(self, data, status=200):
        raw = json.dumps(data).encode()
        self.send_response(status)
        self.send_header("Content-Type", "application/json")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Content-Length", str(len(raw)))
        self.end_headers()
        self.wfile.write(raw)
