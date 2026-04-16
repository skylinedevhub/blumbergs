#!/usr/bin/env python3
"""Generate Blumbergs Snapshot comparison workbook (2023 vs 2024).

Copies the Canada 2024 Summary sheet from the snapshot workbook and builds
a Comparison sheet where every 2024 value is a formula referencing the
Summary, so you can see exactly where each number comes from.

2023 values are hardcoded from the published Blumbergs Canadian Charity
Sector Snapshot 2023 (text highlights and Schedule 6 images).

Usage:
    python3 scripts/reports/generate_comparison.py
"""

import os
import sys
import duckdb
from copy import copy
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, numbers, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import argparse

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DB_PATH = os.path.join(PROJECT_ROOT, "data", "db", "cra_charities.duckdb")
def snapshot_path(year):
    return os.path.join(PROJECT_ROOT, "data", "exports", f"snapshots_{year}", f"snapshot_{year}_canada.xlsx")

def output_path(current_year, prior_year):
    return os.path.join(PROJECT_ROOT, "data", "exports", f"snapshot_comparison_{prior_year}_vs_{current_year}.xlsx")

DATA_SHEET = "Canada 2024"  # name of the embedded source data sheet

# ── Styles ────────────────────────────────────────────────────────────────────
BOLD = Font(bold=True)
BOLD_14 = Font(bold=True, size=14)
BOLD_12 = Font(bold=True, size=12)
HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
PCT_FMT = '0.0%'
NUM_FMT = '#,##0'
MONEY_FMT = '#,##0'

# ── 2023 Published Values ────────────────────────────────────────────────────
# Source: Blumbergs Snapshot of the Canadian Charity Sector 2023 (April 2025)
# "text" = from text highlights (rounded); "sch6" = from Schedule 6 images (exact)

# 2023 exact values from Schedule 6 images in published PDF
SCH6_2023 = {
    "4100": 89_402_068_523,
    "4110": 13_436_483_115,
    "4120": 34_366_146_082,
    "4130": 3_088_825_106,
    "4200": 702_556_461_703,
    "4250": 55_210_860_812,
    "4350": 398_395_161_823,
    "4500": 23_390_731_888,
    "5610": 1_253_983_000,
    "4510": 11_351_302_404,
    "4575_v23": 113_455_262,     # V23: "Tax-receipted from outside Canada"
    "4580_v23": 4_074_915_291,   # V23: "Non-tax-receipted from outside Canada"
    "4590": 3_544_038_905,       # Gross proceeds from disposition of assets
    "4800": 1_486_410_111,
    "4810": 3_629_490_934,
    "4820": 7_258_880_640,
    "4830": 1_920_514_148,
    "4840": 4_834_006_548,
    "4850": 14_655_716_180,
    "4860": 7_847_801_230,
    "4870": 1_039_054_491,
    "4890": 2_029_968_190,
    "4891": 27_960_687_480,
    "4900": 14_348_907_186,
    "4910": 4_953_316_861,
    "5000": 250_713_816_895,
    "5010": 25_677_111_408,
    "5050": 13_797_284_776,
    "5500": 354_027_877,
    "5510": 140_231_072,
    "5750": 67_987_148,
    "5900": 50_384_832_345,
    "5910": 51_484_882_197,
}

# 2023 text-level values (rounded from highlights)
TEXT_2023 = {
    "charity_count": 83_540,
    "active": 78_598,           # "78,598 identified themselves as active"
    "inactive": 3_429,          # "3,429 identified themselves as inactive"
    "gifts_to_qd": 29_814,     # "29,814 made gifts to qualified donees"
    "has_employment": 43_351,   # "43,351 identified having employment expenses"
    "no_employment": 39_775,
    "foreign_activities": 5_089,
    "global_affairs": 139,
    "4700": 393_000_000_000,    # "$393 billion"
    "4540": 13_100_000_000,     # "$13.1 billion"
    "4550": 225_600_000_000,    # "$225.6 billion"
    "4560": 13_000_000_000,     # "$13 billion"
    "4880": 199_000_000_000,    # "$199 billion" (compensation)
    "5100": 354_000_000_000,    # "$354 billion"
    "tax_receipted": 23_400_000_000,  # "$23.4 billion" (rounded)
    "compensation_charities": 43_351,
}


def ref(cell):
    """Return a formula referencing a cell in the Canada 2024 data sheet."""
    return f"='{DATA_SHEET}'!{cell}"


def copy_summary_sheet(wb):
    """Copy the Summary sheet from the Canada 2024 snapshot workbook."""
    print("  Loading Canada 2024 snapshot workbook...")
    src_wb = load_workbook(SNAPSHOT_PATH, data_only=True)
    src_ws = src_wb["Summary"]

    ws = wb.create_sheet(DATA_SHEET)

    # Copy column widths
    for col_letter in ["A", "B", "C", "D", "E"]:
        if col_letter in src_ws.column_dimensions:
            ws.column_dimensions[col_letter].width = src_ws.column_dimensions[col_letter].width

    # Copy cell values and basic formatting
    for row in src_ws.iter_rows(min_row=1, max_row=src_ws.max_row, max_col=src_ws.max_column):
        for cell in row:
            new_cell = ws[cell.coordinate]
            new_cell.value = cell.value
            if cell.font and cell.font.bold:
                new_cell.font = Font(bold=True, size=cell.font.size or 11)
            if cell.number_format and cell.number_format != "General":
                new_cell.number_format = cell.number_format

    print(f"    Copied {src_ws.max_row} rows to '{DATA_SHEET}' sheet")
    src_wb.close()
    return ws


def build_comparison(wb):
    """Build the Comparison sheet with formulas referencing the Canada 2024 sheet.

    The sheet shows 2023 vs 2024 highlight metrics. Every 2024 value is an
    Excel formula pointing to a specific cell in the 'Canada 2024' sheet,
    so users can click any value and see exactly where it came from.
    """
    ws = wb.create_sheet("Comparison", 0)

    # Column widths
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 52
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 45

    # ── Header ──
    ws.append(["Canadian Charity Sector — 2023 vs 2024 Comparison"])
    ws["A1"].font = BOLD_14

    ws.append(["2023: from Blumbergs Snapshot publication (text highlights + Schedule 6 images). "
               "2024: formulas referencing 'Canada 2024' sheet (from CRA T3010 database)."])
    ws.append([])

    # ── Column headers ──
    row = 4
    headers = ["Line", "Metric", "2023 Snapshot", "2024 Database", "Change", "Source / Formula Reference"]
    ws.append(headers)
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_idx)
        cell.font = BOLD
        cell.fill = HEADER_FILL

    # ── Helper to add a data row ──
    def add_row(line, metric, val_2023, ref_2024, source, change_formula=True):
        """Add a comparison row.

        ref_2024: cell address in Canada 2024 sheet (e.g. "C108") or a raw formula string
        """
        nonlocal row
        row += 1
        ws.cell(row=row, column=1, value=line)
        ws.cell(row=row, column=2, value=metric)

        c_cell = ws.cell(row=row, column=3, value=val_2023)
        if isinstance(val_2023, (int, float)) and val_2023 > 1000:
            c_cell.number_format = NUM_FMT

        if ref_2024:
            if ref_2024.startswith("="):
                d_formula = ref_2024
            else:
                d_formula = ref(ref_2024)
            ws.cell(row=row, column=4, value=d_formula)
            ws.cell(row=row, column=4).number_format = NUM_FMT
        else:
            ws.cell(row=row, column=4, value=None)

        if change_formula and val_2023 and ref_2024:
            ws.cell(row=row, column=5,
                    value=f"=IF(C{row}<>0,(D{row}-C{row})/C{row},\"\")")
            ws.cell(row=row, column=5).number_format = PCT_FMT

        ws.cell(row=row, column=6, value=source)
        return row

    def add_section(title):
        nonlocal row
        row += 1
        ws.append([])
        row += 1
        ws.cell(row=row, column=1, value=title)
        ws.cell(row=row, column=1).font = BOLD_12

    def add_blank():
        nonlocal row
        row += 1

    # ── GENERAL ──
    add_row("", "Registered charities (ident count)", TEXT_2023["charity_count"],
            "C7", "2023 text: '83,540' / 2024: 'Canada 2024'!C7")
    add_row("", "  A: Public Foundation", None, "C8",
            "'Canada 2024'!C8", change_formula=False)
    add_row("", "  B: Private Foundation", None, "C9",
            "'Canada 2024'!C9", change_formula=False)
    add_row("", "  C: Charitable Organization", None, "C10",
            "'Canada 2024'!C10", change_formula=False)
    add_blank()
    add_row("1800", "Active charities (C1 Yes)", TEXT_2023["active"],
            "C23", "2023 text / 2024: Section C line 1800 Yes count")
    add_row("1800", "Inactive charities (C1 No)", TEXT_2023["inactive"],
            "D23", "2023 text / 2024: Section C line 1800 No count")
    add_row("2000", "Made gifts to qualified donees (C3 Yes)", TEXT_2023["gifts_to_qd"],
            "C26", "2023 text / 2024: Section C line 2000 Yes count")
    add_row("2100", "Activities outside Canada (C4 Yes)", TEXT_2023["foreign_activities"],
            "C27", "2023 text / 2024: Section C line 2100 Yes count")
    add_row("3400", "Charities with employment expenses (C9 Yes)", TEXT_2023["has_employment"],
            "C59", "2023 text / 2024: Section C line 3400 Yes count")
    add_row("3400", "Charities without employment expenses (C9 No)", TEXT_2023["no_employment"],
            "D59", "2023 text / 2024: Section C line 3400 No count")

    # ── REVENUE ──
    add_section("REVENUE")
    add_row("4700", "Total Revenue", TEXT_2023["4700"],
            "C108", "2023 text: '$393B' / 2024: Section D line 4700")
    add_row("4500", "Tax-Receipted Gifts", SCH6_2023["4500"],
            "C95", "2023 Sch6 exact / 2024: Section D line 4500")
    add_row("5610", "Tax-Receipted Tuition Fees", SCH6_2023["5610"],
            "C96", "2023 Sch6 exact / 2024: Section D line 5610")
    add_row("4510", "Gifts from other registered charities", SCH6_2023["4510"],
            "C97", "2023 Sch6 exact / 2024: Section D line 4510")
    add_row("4540", "Revenue from FEDERAL government", TEXT_2023["4540"],
            "C99", "2023 text: '$13.1B' / 2024: Section D line 4540")
    add_row("4550", "Revenue from PROVINCIAL/TERRITORIAL governments", TEXT_2023["4550"],
            "C100", "2023 text: '$225.6B' / 2024: Section D line 4550")
    add_row("4560", "Revenue from MUNICIPAL/REGIONAL governments", TEXT_2023["4560"],
            "C101", "2023 text: '$13B' / 2024: Section D line 4560")

    # Computed total government
    govt_2023 = TEXT_2023["4540"] + TEXT_2023["4550"] + TEXT_2023["4560"]
    r_govt = add_row("calc", "Total Government (4540+4550+4560)", govt_2023,
                     "C102", "2023: sum of text values / 2024: computed in Section D")
    r_rev = row - 7  # row of Total Revenue (4700) — adjust if layout changes

    # Government % of revenue
    add_row("calc", "Government as % of Total Revenue", 0.64,
            f"=D{r_govt}/D{r_rev}", "2023 text: '64%' / 2024: computed D{}/D{}".format(r_govt, r_rev),
            change_formula=False)
    ws.cell(row=row, column=3).number_format = '0.0%'
    ws.cell(row=row, column=4).number_format = '0.0%'

    add_row("4575", "Revenue from outside Canada (non-tax-receipted) [V24 definition]",
            None, "C104",
            "2024: Section D line 4575. Line definition changed between V23→V24.",
            change_formula=False)
    add_row("4575", "Revenue from outside Canada (tax-receipted) [V23 definition]",
            SCH6_2023["4575_v23"], None,
            "2023 Sch6: V23 line 4575 = tax-receipted from outside Canada. No V24 equivalent.",
            change_formula=False)
    add_row("4571", "Tax-receipted revenue from outside Canada [V24 line]",
            None, "C103",
            "2024: Section D line 4571 (new in V24). No 2023 equivalent.",
            change_formula=False)
    add_row("4580", "Interest/investment income [V24 definition]",
            None, "='" + DATA_SHEET + "'!C229",
            "2024: Sch6 line 4580. Line definition changed between V23→V24.",
            change_formula=False)
    add_row("4580", "Non-tax-receipted from outside Canada [V23 definition]",
            SCH6_2023["4580_v23"], None,
            "2023 Sch6: V23 line 4580 = non-tax-receipted from outside Canada.",
            change_formula=False)

    # ── EXPENDITURES ──
    add_section("EXPENDITURES")
    add_row("5100", "Total Expenditures", TEXT_2023["5100"],
            "C119", "2023 text: '$354B' / 2024: Section D line 5100")
    add_row("5000", "  Charitable Activities", SCH6_2023["5000"],
            "C115", "2023 Sch6 exact / 2024: Section D line 5000")
    add_row("5010", "  Management and Administration", SCH6_2023["5010"],
            "C116", "2023 Sch6 exact / 2024: Section D line 5010")
    add_row("5020", "  Fundraising", None,
            "='" + DATA_SHEET + "'!C257",
            "2024: Sch6 line 5020. No 2023 Sch6 value available.",
            change_formula=False)
    add_row("5045", "Grants to non-qualified donees", None,
            "C117", "2024: Section D line 5045. No 2023 Sch6 value.",
            change_formula=False)
    add_row("5050", "Gifts to Qualified Donees", SCH6_2023["5050"],
            "C118", "2023 Sch6 exact / 2024: Section D line 5050")
    add_row("4880", "Total Compensation (financial_d line 4880)", TEXT_2023["4880"],
            "='" + DATA_SHEET + "'!C248",
            "2023 text: '$199B' / 2024: Sch6 line 4880")
    add_row("390", "Total Compensation (Schedule 3 line 390)", None,
            "C162", "2024: Schedule 3 line 390. Cross-check with 4880.",
            change_formula=False)

    # ── BALANCE SHEET ──
    add_section("BALANCE SHEET")
    add_row("4200", "Total Assets", SCH6_2023["4200"],
            "C88", "2023 Sch6 exact / 2024: Section D line 4200")
    add_row("4350", "Total Liabilities", SCH6_2023["4350"],
            "C89", "2023 Sch6 exact / 2024: Section D line 4350")
    add_row("4100", "Cash and Short-Term Investments", SCH6_2023["4100"],
            "='" + DATA_SHEET + "'!C190",
            "2023 Sch6 exact / 2024: Sch6 line 4100")
    add_row("4155", "Land and Buildings in Canada", None,
            "='" + DATA_SHEET + "'!C198",
            "2024: Sch6 line 4155. No 2023 Sch6 value.",
            change_formula=False)

    # ── EMPLOYMENT ──
    add_section("EMPLOYMENT")
    add_row("300", "Full-time positions (Schedule 3)", None,
            "C148", "2024: Schedule 3 line 300.",
            change_formula=False)
    add_row("370", "Part-time/seasonal positions (Schedule 3)", None,
            "C160", "2024: Schedule 3 line 370.",
            change_formula=False)

    # ── FOREIGN ACTIVITIES ──
    add_section("FOREIGN ACTIVITIES (Schedule 2)")
    add_row("200", "Total expenditures outside Canada", None,
            "C136", "2024: Schedule 2 line 200.",
            change_formula=False)
    add_row("220", "Projects funded by Global Affairs Canada (Yes)", TEXT_2023["global_affairs"],
            "C138", "2023 text: '139' / 2024: Schedule 2 line 220 Yes count")

    # ── DAF ──
    add_section("DONOR ADVISED FUNDS")
    add_row("5860", "Charities holding DAF accounts (Yes)", None,
            "C74", "2024: Section C line 5860 Yes count.",
            change_formula=False)
    add_row("5861", "Total DAF accounts", None,
            "C75", "2024: Section C line 5861.",
            change_formula=False)
    add_row("5862", "Total value of DAF accounts", None,
            "C76", "2024: Section C line 5862.",
            change_formula=False)
    add_row("5863", "Donations to DAFs received", None,
            "C77", "2024: Section C line 5863.",
            change_formula=False)
    add_row("5864", "Qualifying disbursements from DAFs", None,
            "C78", "2024: Section C line 5864.",
            change_formula=False)

    # ── NOTES ──
    add_blank()
    add_blank()
    row += 1
    ws.cell(row=row, column=1, value="NOTES:").font = BOLD
    notes = [
        "1. 2023 values from Blumbergs Snapshot of the Canadian Charity Sector 2023 (April 2025).",
        "2. 'text' values are rounded from the publication highlights (e.g. '$393 billion' → 393,000,000,000).",
        "3. 'Sch6' values are exact numbers read from Schedule 6 images in the published PDF.",
        "4. 2024 values are Excel formulas referencing the 'Canada 2024' sheet — click any cell to see the source.",
        "5. Lines 4575, 4580 changed definition between V23 and V24 forms — comparisons are NOT apples-to-apples.",
        "6. Line 4570 (total govt) is unreliable; total government always computed as 4540+4550+4560.",
        "7. 2023 covers ~83,540 charities filed; 2024 has 83,093 in financial_d (83,275 in ident).",
        "8. Line 4101/4102 sub-breakdown of 4100 changed between V23→V24 (amounts receivable → cash/investments).",
    ]
    for note in notes:
        row += 1
        ws.cell(row=row, column=1, value=note)

    print(f"    Comparison sheet: {row} rows")
    return ws


def build_all_financial_lines(wb):
    """Build the 'All Financial Lines' sheet with 2023 values and 2024 formulas.

    2024 values reference the Schedule 6 section of the Canada 2024 sheet
    (rows 190-268 of the Summary).
    """
    ws = wb.create_sheet("All Financial Lines")

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 40

    ws.append([f"T3010 Section D — All Line Totals (2023 vs 2024)"])
    ws["A1"].font = BOLD_14
    ws.append(["2023: from Blumbergs Snapshot Sch6 images. "
               "2024: formulas → 'Canada 2024' sheet Schedule 6 section."])
    ws.append([])

    headers = ["Line #", "Description", "2023 Snapshot", "2024 Database", "Change %", "Notes"]
    ws.append(headers)
    for col_idx in range(1, 7):
        ws.cell(row=4, column=col_idx).font = BOLD
        ws.cell(row=4, column=col_idx).fill = HEADER_FILL

    # Schedule 6 line → row in Canada 2024 Summary sheet
    # These row numbers are deterministic from generate_snapshot.py
    SCH6_ROW = {
        "4100": 190, "4101": 191, "4102": 192, "4110": 193, "4120": 194,
        "4130": 195, "4140": 196, "4150": 197, "4155": 198, "4157": 199,
        "4158": 200, "4160": 201, "4165": 202, "4166": 203, "4170": 204,
        "4190": 205, "4200": 206,
        "4300": 209, "4310": 210, "4320": 211, "4330": 212, "4350": 213,
        "4250": 215,
        "4500": 218, "5610": 219, "4510": 220, "4530": 221,
        "4540": 222, "4550": 223, "4560": 224,
        "4571": 225, "4575": 226, "4576": 227, "4577": 228, "4580": 229,
        "4590": 230, "4600": 231, "4610": 232, "4620": 233,
        "4630": 234, "4640": 235, "4650": 236, "4700": 237,
        "4800": 240, "4810": 241, "4820": 242, "4830": 243,
        "4840": 244, "4850": 245, "4860": 246, "4870": 247,
        "4880": 248, "4890": 249, "4891": 250, "4900": 251,
        "4910": 252, "4920": 253, "4950": 254,
        "5000": 255, "5010": 256, "5020": 257, "5040": 258,
        "5045": 259, "5050": 260, "5100": 261,
        "5500": 264, "5510": 265, "5750": 266, "5900": 267, "5910": 268,
    }

    # Lines to include: (line_num, description, 2023_value_or_None, notes)
    # None for 2023 means no published value available
    lines = [
        ("ASSETS",),
        ("4100", "Cash, bank accounts, short-term investments", SCH6_2023.get("4100"), ""),
        ("4101", "  Cash and bank accounts [V24] / Receivables non-arm's [V23]",
         None, "Sub-line definition changed V23→V24"),
        ("4102", "  Short-term investments [V24] / Receivables others [V23]",
         None, "Sub-line definition changed V23→V24"),
        ("4110", "Amounts receivable from non-arm's length persons", SCH6_2023.get("4110"), ""),
        ("4120", "Amounts receivable from all others", SCH6_2023.get("4120"), ""),
        ("4130", "Investments in non-arm's length persons", SCH6_2023.get("4130"), ""),
        ("4140", "Long-term investments", None, ""),
        ("4150", "Inventories", None, ""),
        ("4155", "Land and buildings in Canada", None, ""),
        ("4157", "  Used for charitable programs or admin", None, ""),
        ("4158", "  Used for other purposes", None, ""),
        ("4160", "Other capital assets in Canada", None, ""),
        ("4165", "Capital assets outside Canada", None, ""),
        ("4166", "Accumulated amortization of capital assets", None, ""),
        ("4170", "Other assets", None, ""),
        ("4190", "Impact investments", None, ""),
        ("4200", "TOTAL ASSETS", SCH6_2023.get("4200"), ""),
        ("",),
        ("LIABILITIES",),
        ("4300", "Accounts payable and accrued liabilities", None, ""),
        ("4310", "Deferred revenue", None, ""),
        ("4320", "Amounts owing to non-arm's length persons", None, ""),
        ("4330", "Other liabilities", None, ""),
        ("4350", "TOTAL LIABILITIES", SCH6_2023.get("4350"), ""),
        ("",),
        ("4250", "Property not used in charitable activities", SCH6_2023.get("4250"), ""),
        ("",),
        ("REVENUE",),
        ("4500", "Tax-receipted gifts (donation receipts issued)", SCH6_2023.get("4500"), "Sch6 exact"),
        ("5610", "Tax-receipted tuition fees", SCH6_2023.get("5610"), "Sch6 exact"),
        ("4510", "Gifts from other registered charities", SCH6_2023.get("4510"), "Sch6 exact"),
        ("4530", "Other gifts (no tax receipt)", None, ""),
        ("4540", "Revenue from FEDERAL government", TEXT_2023["4540"], "Text: $13.1B"),
        ("4550", "Revenue from PROVINCIAL/TERRITORIAL governments", TEXT_2023["4550"], "Text: $225.6B"),
        ("4560", "Revenue from MUNICIPAL/REGIONAL governments", TEXT_2023["4560"], "Text: $13B"),
        ("4571", "Tax-receipted revenue from outside Canada (govt+non-govt) [V24]", None,
         "New line in V24"),
        ("4575", "Non-tax-receipted revenue from outside Canada [V24 defn]", None,
         "V23 4575 = tax-receipted from outside Canada (different field!)"),
        ("4576", "Interest/investment from impact investments", None, "New in V24"),
        ("4577", "Interest/investment from non-arm's length persons", None, "New in V24"),
        ("4580", "Interest/investment income received or earned [V24 defn]", None,
         "V23 4580 = non-tax-receipted from outside Canada (different field!)"),
        ("4590", "Gross proceeds from disposition of assets", SCH6_2023.get("4590"), "Sch6 exact"),
        ("4600", "Net proceeds from disposition of assets", None, ""),
        ("4610", "Gross income from rental of land/buildings", None, ""),
        ("4620", "Membership fees, dues, association fees", None, ""),
        ("4630", "Non-tax-receipted revenue from fundraising", None, ""),
        ("4640", "Revenue from sale of goods and services", None, ""),
        ("4650", "Other revenue", None, ""),
        ("4700", "TOTAL REVENUE", TEXT_2023["4700"], "Text: $393B"),
        ("",),
        ("EXPENDITURES",),
        ("4800", "Advertising and promotion", SCH6_2023.get("4800"), "Sch6 exact"),
        ("4810", "Travel and vehicle expenses", SCH6_2023.get("4810"), "Sch6 exact"),
        ("4820", "Interest and bank charges", SCH6_2023.get("4820"), "Sch6 exact"),
        ("4830", "Licences, memberships, and dues", SCH6_2023.get("4830"), "Sch6 exact"),
        ("4840", "Office supplies and expenses", SCH6_2023.get("4840"), "Sch6 exact"),
        ("4850", "Occupancy costs", SCH6_2023.get("4850"), "Sch6 exact"),
        ("4860", "Professional and consulting fees", SCH6_2023.get("4860"), "Sch6 exact"),
        ("4870", "Education and training for staff and volunteers", SCH6_2023.get("4870"), "Sch6 exact"),
        ("4880", "Total expenditure on all compensation", TEXT_2023["4880"], "Text: $199B"),
        ("4890", "Fair market value of donated goods used", SCH6_2023.get("4890"), "Sch6 exact"),
        ("4891", "Purchased supplies and assets", SCH6_2023.get("4891"), "Sch6 exact"),
        ("4900", "Amortization of capitalized assets", SCH6_2023.get("4900"), "Sch6 exact"),
        ("4910", "Research grants and scholarships", SCH6_2023.get("4910"), "Sch6 exact"),
        ("4920", "All other expenditures", None, ""),
        ("4950", "Total expenditures before qualifying disbursements", None, ""),
        ("5000", "  (a) Charitable activities", SCH6_2023.get("5000"), "Sch6 exact"),
        ("5010", "  (b) Management and administration", SCH6_2023.get("5010"), "Sch6 exact"),
        ("5020", "  (c) Fundraising", None, ""),
        ("5040", "  (d) Other expenditures included in 4950", None, ""),
        ("5045", "Grants to non-qualified donees", None, ""),
        ("5050", "Gifts to all qualified donees", SCH6_2023.get("5050"), "Sch6 exact"),
        ("5100", "TOTAL EXPENDITURES", TEXT_2023["5100"], "Text: $354B"),
        ("",),
        ("OTHER FINANCIAL INFORMATION",),
        ("5500", "Amount accumulated (permission to accumulate)", SCH6_2023.get("5500"), "Sch6 exact"),
        ("5510", "Amount disbursed for specified purpose", SCH6_2023.get("5510"), "Sch6 exact"),
        ("5750", "Permission to reduce disbursement quota", SCH6_2023.get("5750"), "Sch6 exact"),
        ("5900", "Property not used — beginning of period", SCH6_2023.get("5900"), "Sch6 exact"),
        ("5910", "Property not used — end of period", SCH6_2023.get("5910"), "Sch6 exact"),
    ]

    r = 4
    for entry in lines:
        r += 1
        if len(entry) == 1:
            # Section header or blank
            if entry[0]:
                ws.cell(row=r, column=1, value=entry[0]).font = BOLD_12
            continue

        line_num, desc, val_2023, notes = entry

        ws.cell(row=r, column=1, value=line_num)
        ws.cell(row=r, column=2, value=desc)

        if val_2023 is not None:
            c = ws.cell(row=r, column=3, value=val_2023)
            c.number_format = NUM_FMT

        # 2024 formula referencing Schedule 6 section
        if line_num in SCH6_ROW:
            summary_row = SCH6_ROW[line_num]
            ws.cell(row=r, column=4, value=ref(f"C{summary_row}"))
            ws.cell(row=r, column=4).number_format = NUM_FMT

            # Change % formula
            if val_2023 is not None:
                ws.cell(row=r, column=5, value=f"=IF(C{r}<>0,(D{r}-C{r})/C{r},\"\")")
                ws.cell(row=r, column=5).number_format = PCT_FMT

        ws.cell(row=r, column=6, value=notes)

    print(f"    All Financial Lines sheet: {r} rows")
    return ws


def build_by_designation(wb):
    """Build the By Designation sheet with 2024 data from database."""
    con = duckdb.connect(DB_PATH, read_only=True)
    ws = wb.create_sheet("By Designation")

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 22

    ws.append(["2024 Key Totals by Designation"])
    ws["A1"].font = BOLD_14
    ws.append([])
    headers = ["Designation", "Count", "Total Revenue", "Total Expenditures", "Total Assets"]
    ws.append(headers)
    for col_idx in range(1, 6):
        ws.cell(row=3, column=col_idx).font = BOLD
        ws.cell(row=3, column=col_idx).fill = HEADER_FILL

    def money(col):
        return f"TRY_CAST(REPLACE(REPLACE(t.\"{col}\", '$', ''), ',', '') AS DECIMAL)"

    designations = [
        ("A", "A — Public Foundation"),
        ("B", "B — Private Foundation"),
        ("C", "C — Charitable Organization"),
    ]

    total_count = 0
    total_rev = 0
    total_exp = 0
    total_assets = 0

    for code, label in designations:
        row = con.execute(f"""
            SELECT
                COUNT(*) AS cnt,
                SUM({money('4700')}) AS rev,
                SUM({money('5100')}) AS exp,
                SUM({money('4200')}) AS assets
            FROM financial_d t
            INNER JOIN charity_base cb ON t."BN/Registration Number" = cb.bn
                AND t.data_year = cb.data_year
            WHERE cb.designation_code = '{code}'
                AND cb.data_year = (SELECT MAX(data_year) FROM charity_base)
        """).fetchone()
        cnt, rev, exp, assets = row
        ws.append([label, cnt, rev, exp, assets])
        for col_idx in range(2, 6):
            ws.cell(row=ws.max_row, column=col_idx).number_format = NUM_FMT
        total_count += cnt or 0
        total_rev += rev or 0
        total_exp += exp or 0
        total_assets += assets or 0

    ws.append(["TOTAL", total_count, total_rev, total_exp, total_assets])
    for col_idx in range(1, 6):
        ws.cell(row=ws.max_row, column=col_idx).font = BOLD
        ws.cell(row=ws.max_row, column=col_idx).number_format = NUM_FMT

    con.close()
    print("    By Designation sheet: done")
    return ws


def build_compensation(wb):
    """Build the Compensation sheet with 2023 vs 2024 comparison."""
    ws = wb.create_sheet("Compensation")

    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 12

    ws.append(["Schedule 3 Compensation — 2023 vs 2024"])
    ws["A1"].font = BOLD_14
    ws.append([])
    headers = ["Metric", "2023 Snapshot", "2024 Database", "Change"]
    ws.append(headers)
    for col_idx in range(1, 5):
        ws.cell(row=3, column=col_idx).font = BOLD
        ws.cell(row=3, column=col_idx).fill = HEADER_FILL

    # Sch3 row references in Canada 2024 Summary:
    # Row 148: 300 FT positions
    # Row 160: 370 PT positions
    # Row 161: 380 PT compensation
    # Row 162: 390 Total compensation
    # Row 164: Cross-check 4880

    rows = [
        ("Charities with compensation data", TEXT_2023["compensation_charities"], "C59",
         "Section C line 3400 Yes count"),
        ("Full-time positions (line 300)", None, "C148", "Schedule 3 line 300"),
        ("Part-time positions (line 370)", None, "C160", "Schedule 3 line 370"),
        ("Total compensation (line 390)", TEXT_2023["4880"], "C162", "Schedule 3 line 390"),
        ("Cross-check: financial_d line 4880", None, "C164", "Cross-check value"),
    ]

    r = 3
    for metric, val_2023, ref_cell, _note in rows:
        r += 1
        ws.cell(row=r, column=1, value=metric)
        if val_2023 is not None:
            ws.cell(row=r, column=2, value=val_2023)
            ws.cell(row=r, column=2).number_format = NUM_FMT
        ws.cell(row=r, column=3, value=ref(ref_cell))
        ws.cell(row=r, column=3).number_format = NUM_FMT
        if val_2023 is not None:
            ws.cell(row=r, column=4, value=f"=IF(B{r}<>0,(C{r}-B{r})/B{r},\"\")")
            ws.cell(row=r, column=4).number_format = PCT_FMT

    # Salary band breakdown (2024 only, from Schedule 3)
    r += 2
    ws.cell(row=r, column=1, value="SALARY BANDS (Full-time, 2024)").font = BOLD
    bands = [
        ("$1 - $39,999 (line 305)", "C149"),
        ("$40,000 - $79,999 (line 310)", "C150"),
        ("$80,000 - $119,999 (line 315)", "C151"),
        ("$120,000 - $159,999 (line 320)", "C152"),
        ("$160,000 - $199,999 (line 325)", "C153"),
        ("$200,000 - $249,999 (line 330)", "C154"),
        ("$250,000 - $299,999 (line 335)", "C155"),
        ("$300,000 - $349,999 (line 340)", "C156"),
        ("$350,000 and over (line 345)", "C157"),
    ]
    for band_label, band_ref in bands:
        r += 1
        ws.cell(row=r, column=1, value=band_label)
        ws.cell(row=r, column=3, value=ref(band_ref))
        ws.cell(row=r, column=3).number_format = NUM_FMT

    print("    Compensation sheet: done")
    return ws


def main():
    parser = argparse.ArgumentParser(description="Generate Blumbergs Snapshot comparison workbook")
    parser.add_argument("--year", type=int, default=None,
                        help="Current year for comparison (default: latest in database)")
    args = parser.parse_args()

    con = duckdb.connect(DB_PATH, read_only=True)
    if args.year:
        current_year = args.year
    else:
        current_year = con.execute("SELECT MAX(data_year) FROM charity_base").fetchone()[0]
    prior_year = current_year - 1
    con.close()

    sp = snapshot_path(current_year)
    op = output_path(current_year, prior_year)

    if not os.path.exists(sp):
        print(f"ERROR: Canada {current_year} snapshot not found: {sp}")
        print(f"Run: python3 scripts/reports/generate_snapshot.py --year {current_year}")
        sys.exit(1)

    print(f"Generating comparison workbook ({prior_year} vs {current_year})...")
    wb = Workbook()

    # 1. Copy the Canada snapshot Summary as a source data sheet
    # Need to temporarily set module-level SNAPSHOT_PATH for copy_summary_sheet
    global SNAPSHOT_PATH, OUTPUT_PATH, DATA_SHEET
    SNAPSHOT_PATH = sp
    OUTPUT_PATH = op
    DATA_SHEET = f"Canada {current_year}"

    copy_summary_sheet(wb)

    # 2. Build comparison sheets
    print("  Building comparison sheets...")
    build_comparison(wb)
    build_all_financial_lines(wb)
    build_by_designation(wb)
    build_compensation(wb)

    # Remove default empty sheet
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
        del wb["Sheet"]

    # Ensure Comparison is the first (active) sheet
    wb.move_sheet("Comparison", offset=-len(wb.sheetnames) + 1)

    os.makedirs(os.path.dirname(op), exist_ok=True)
    wb.save(op)
    print(f"\nSaved to: {op}")
    print(f"Sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
