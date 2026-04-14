#!/usr/bin/env python3
"""Replace embedded T3010 PDF images in existing snapshot article Word documents.

Opens each .docx in the project root, removes old T3010 page images,
re-renders the corresponding PDF from snapshot_pdfs_2024/ at full-page size,
and inserts the new images filling the printable area. Saves as v2.

Usage:
    python3 scripts/reports/update_article_pdfs.py          # Update all 13 articles
    python3 scripts/reports/update_article_pdfs.py --dry-run # Preview mapping only
"""

import argparse
import io
import os
import sys

import fitz  # PyMuPDF
from docx import Document
from docx.shared import Inches, Emu, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PDF_DIR = os.path.join(PROJECT_ROOT, "data", "exports", "snapshot_pdfs_2024")

# Map each article .docx filename to its corresponding PDF filename
DOCX_TO_PDF = {
    "Blumbergs-Snapshot-Canadian-Charity-Sector-2024.docx": "Blumbergs-Snapshot-T3010-2024-canada.pdf",
    "Blumbergs-Snapshot-Ontario-Charity-Sector-2024.docx": "Blumbergs-Snapshot-T3010-2024-ON.pdf",
    "Blumbergs-Snapshot-Quebec-Charity-Sector-2024.docx": "Blumbergs-Snapshot-T3010-2024-QC.pdf",
    "Blumbergs-Snapshot-British-Columbia-Charity-Sector-2024.docx": "Blumbergs-Snapshot-T3010-2024-BC.pdf",
    "Blumbergs-Snapshot-Alberta-Charity-Sector-2024.docx": "Blumbergs-Snapshot-T3010-2024-AB.pdf",
    "Blumbergs-Snapshot-Manitoba-Charity-Sector-2024.docx": "Blumbergs-Snapshot-T3010-2024-MB.pdf",
    "Blumbergs-Snapshot-Saskatchewan-Charity-Sector-2024.docx": "Blumbergs-Snapshot-T3010-2024-SK.pdf",
    "Blumbergs-Snapshot-Nova-Scotia-Charity-Sector-2024.docx": "Blumbergs-Snapshot-T3010-2024-NS.pdf",
    "Blumbergs-Snapshot-New-Brunswick-Charity-Sector-2024.docx": "Blumbergs-Snapshot-T3010-2024-NB.pdf",
    "Blumbergs-Snapshot-Atlantic-Provinces-Charity-Sector-2024.docx": "Blumbergs-Snapshot-T3010-2024-atlantic.pdf",
    "Blumbergs-Snapshot-Public-Foundations-in-the-Canadian-Charity-Sector-2024.docx": "Blumbergs-Snapshot-T3010-2024-designation_A.pdf",
    "Blumbergs-Snapshot-Private-Foundations-in-the-Canadian-Charity-Sector-2024.docx": "Blumbergs-Snapshot-T3010-2024-designation_B.pdf",
    "Blumbergs-Snapshot-Charitable-Organizations-in-the-Canadian-Charity-Sector-2024.docx": "Blumbergs-Snapshot-T3010-2024-designation_C.pdf",
}

WML_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"


def elem_text(elem):
    """Extract concatenated text from a <w:p> element."""
    return "".join(t.text or "" for t in elem.findall(f".//{{{WML_NS}}}t")).strip()


def has_image(elem):
    """Check if a <w:p> element contains an inline drawing (image)."""
    return len(elem.findall(f".//{{{WP_NS}}}inline")) > 0


def render_pdf_pages(pdf_path, dpi=250):
    """Render all pages of a PDF to PNG byte buffers at high DPI."""
    doc = fitz.open(pdf_path)
    pages = []
    for page in doc:
        pix = page.get_pixmap(dpi=dpi)
        pages.append((pix.tobytes("png"), page.rect.width / 72, page.rect.height / 72))
    doc.close()
    return pages


def set_para_spacing_zero(paragraph):
    """Set paragraph before/after spacing to zero. Leave line spacing as auto
    so Word expands the line to fit inline images."""
    pPr = paragraph._element.get_or_add_pPr()
    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        pPr.append(spacing)
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")


def set_page_break_before(paragraph):
    """Set the 'page break before' property on a paragraph."""
    pPr = paragraph._element.get_or_add_pPr()
    pb = OxmlElement("w:pageBreakBefore")
    pPr.append(pb)


def update_article(docx_path, output_path, pdf_path):
    """Replace T3010 images in a Word document with full-page PDF renderings."""
    doc = Document(docx_path)
    body = doc.element.body
    elements = list(body)

    # Read page dimensions from the document section
    section = doc.sections[0]
    avail_w_in = section.page_width.inches - section.left_margin.inches - section.right_margin.inches
    avail_h_in = section.page_height.inches - section.top_margin.inches - section.bottom_margin.inches

    # Find "Further information" paragraph
    further_idx = None
    for i, elem in enumerate(elements):
        tag = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag
        if tag == "p" and "Further information" in elem_text(elem):
            further_idx = i
            break

    if further_idx is None:
        print("    WARNING: Could not find 'Further information' paragraph, skipping")
        return False

    # Walk backward from "Further information" and collect paragraphs to remove
    to_remove = []
    for i in range(further_idx - 1, -1, -1):
        elem = elements[i]
        tag = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag
        if tag != "p":
            break
        text = elem_text(elem)
        is_img = has_image(elem)
        if is_img:
            to_remove.append(elem)
        elif not text:
            to_remove.append(elem)
        else:
            break

    removed_count = len(to_remove)
    for elem in to_remove:
        body.remove(elem)
    print(f"    Removed {removed_count} old image/spacer paragraphs")

    # Re-find "Further information" after removals
    further_elem = None
    for elem in list(body):
        tag = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag
        if tag == "p" and "Further information" in elem_text(elem):
            further_elem = elem
            break

    if further_elem is None:
        print("    WARNING: Lost 'Further information' anchor after removal")
        return False

    # Render new PDF pages at higher DPI for crisp full-page images
    print(f"    Rendering PDF pages from {os.path.basename(pdf_path)}...")
    page_data = render_pdf_pages(pdf_path, dpi=250)
    print(f"    Inserting {len(page_data)} T3010 pages (full-page @ {avail_w_in:.1f}\" x {avail_h_in:.1f}\")...")

    for i, (png_bytes, pdf_w_in, pdf_h_in) in enumerate(page_data):
        # Scale to fit the printable area: constrain by width or height
        scale_w = avail_w_in / pdf_w_in
        scale_h = avail_h_in / pdf_h_in
        scale = min(scale_w, scale_h)
        img_w = pdf_w_in * scale
        img_h = pdf_h_in * scale

        new_para = doc.add_paragraph()

        # Page break before each image page (except first — the previous
        # text content ends the prior page naturally)
        if i > 0:
            set_page_break_before(new_para)

        # Zero out paragraph spacing so the image fills the page
        set_para_spacing_zero(new_para)

        # Insert the image at computed full-page dimensions
        run = new_para.add_run()
        run.add_picture(io.BytesIO(png_bytes), width=Inches(img_w), height=Inches(img_h))

        # Move paragraph before "Further information"
        further_elem.addprevious(new_para._element)

    # Add a page-break paragraph before "Further information" so it starts
    # on a fresh page after the T3010 images
    break_para = doc.add_paragraph()
    set_page_break_before(break_para)
    further_elem.addprevious(break_para._element)

    doc.save(output_path)
    print(f"    Saved: {output_path}")
    return True


def v2_name(docx_name):
    """Insert 'v2' before the .docx extension."""
    base, ext = os.path.splitext(docx_name)
    return f"{base}-v2{ext}"


def main():
    parser = argparse.ArgumentParser(description="Replace T3010 PDF images in snapshot articles (v2)")
    parser.add_argument("--dry-run", action="store_true", help="Show mapping without modifying files")
    args = parser.parse_args()

    updated = 0
    skipped = 0

    for docx_name, pdf_name in sorted(DOCX_TO_PDF.items()):
        docx_path = os.path.join(PROJECT_ROOT, docx_name)
        pdf_path = os.path.join(PDF_DIR, pdf_name)
        out_name = v2_name(docx_name)
        out_path = os.path.join(PROJECT_ROOT, out_name)

        if not os.path.exists(docx_path):
            print(f"SKIP: {docx_name} (not found)")
            skipped += 1
            continue
        if not os.path.exists(pdf_path):
            print(f"SKIP: {docx_name} -> {pdf_name} (PDF not found)")
            skipped += 1
            continue

        print(f"\n{docx_name}  ->  {out_name}")
        print(f"  <- {pdf_name}")

        if args.dry_run:
            updated += 1
            continue

        if update_article(docx_path, out_path, pdf_path):
            updated += 1
        else:
            skipped += 1

    print(f"\nDone: {updated} updated, {skipped} skipped")


if __name__ == "__main__":
    main()
