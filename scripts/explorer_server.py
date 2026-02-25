#!/usr/bin/env python3
"""T3010 Data Explorer — local query server.

Usage:
    python3 scripts/explorer_server.py          # default port 8765
    python3 scripts/explorer_server.py 9000     # custom port

Opens http://localhost:<port> in your browser.
"""

import http.server
import json
import duckdb
import re
import sys
import io
import webbrowser
import threading
import decimal
import datetime
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DB = ROOT / "data" / "db" / "cra_charities.duckdb"
PORT = int(sys.argv[1]) if len(sys.argv) > 1 else 8765


def connect():
    return duckdb.connect(str(DB), read_only=True)


def prep(sql):
    """Strip money() macro definition; inline money() calls."""
    sql = re.sub(
        r"(--[^\n]*\n\s*)?CREATE\s+OR\s+REPLACE\s+MACRO\s+money.*?;\s*", "", sql, flags=re.S | re.I
    )
    sql = re.sub(
        r"\bmoney\(([^)]+)\)",
        r"TRY_CAST(REPLACE(REPLACE(\1, '$', ''), ',', '') AS DECIMAL)",
        sql,
    )
    return sql.strip()


def only_select(sql):
    # Strip SQL comments before checking
    s = re.sub(r"--[^\n]*", "", sql).strip().upper()
    if not (s.startswith("SELECT") or s.startswith("WITH")):
        return False
    for kw in (
        "INSERT", "UPDATE", "DELETE", "DROP", "ALTER",
        "CREATE", "TRUNCATE", "COPY", "ATTACH", "DETACH",
    ):
        if re.search(r"\b" + kw + r"\b", s):
            return False
    return True


def default_ser(o):
    if isinstance(o, decimal.Decimal):
        return float(o)
    if isinstance(o, (datetime.date, datetime.datetime)):
        return o.isoformat()
    return str(o)


class Handler(http.server.SimpleHTTPRequestHandler):
    def __init__(self, *a, **kw):
        super().__init__(*a, directory=str(ROOT), **kw)

    def do_GET(self):
        if self.path in ("/", "/index.html"):
            self.path = "/data-explorer.html"
        elif self.path == "/api/lookups":
            return self._lookups()
        super().do_GET()

    def do_POST(self):
        body = self.rfile.read(int(self.headers.get("Content-Length", 0))).decode()
        if self.path == "/api/query":
            self._query(body)
        elif self.path == "/api/export":
            self._export(body)
        else:
            self.send_error(404)

    def _query(self, body):
        try:
            data = json.loads(body)
            sql = prep(data.get("sql", ""))
            if not only_select(sql):
                return self._json({"error": "Only SELECT queries allowed"}, 400)
            t0 = time.perf_counter()
            con = connect()
            res = con.execute(sql)
            cols = [d[0] for d in res.description]
            rows = [list(r) for r in res.fetchall()]
            elapsed = time.perf_counter() - t0
            con.close()
            self._json(
                {"columns": cols, "rows": rows, "count": len(rows), "time": round(elapsed, 4)}
            )
        except Exception as e:
            self._json({"error": str(e)}, 500)

    def _export(self, body):
        try:
            from openpyxl import Workbook
            from openpyxl.utils import get_column_letter
            from openpyxl.styles import Font

            data = json.loads(body)
            sql = prep(data.get("sql", ""))
            sql = re.sub(r"\bLIMIT\s+\d+\s*;?\s*$", ";", sql, flags=re.I)
            if not only_select(sql):
                return self._json({"error": "Only SELECT queries allowed"}, 400)

            con = connect()
            res = con.execute(sql)
            cols = [d[0] for d in res.description]
            rows = res.fetchall()
            con.close()

            wb = Workbook()
            ws = wb.active
            ws.title = "T3010 Export"
            hfont = Font(bold=True)
            for ci, col in enumerate(cols, 1):
                cell = ws.cell(row=1, column=ci, value=col)
                cell.font = hfont

            for ri, row in enumerate(rows, 2):
                for ci, val in enumerate(row, 1):
                    if val is None:
                        continue
                    if isinstance(val, decimal.Decimal):
                        val = float(val)
                    ws.cell(row=ri, column=ci, value=val)

            for ci in range(1, len(cols) + 1):
                sample = [len(str(cols[ci - 1]))]
                for ri in range(min(50, len(rows))):
                    sample.append(len(str(rows[ri][ci - 1] or "")))
                ws.column_dimensions[get_column_letter(ci)].width = min(max(sample) + 2, 45)

            buf = io.BytesIO()
            wb.save(buf)
            raw = buf.getvalue()

            self.send_response(200)
            self.send_header(
                "Content-Type",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            self.send_header(
                "Content-Disposition", 'attachment; filename="t3010_export.xlsx"'
            )
            self.send_header("Content-Length", str(len(raw)))
            self.end_headers()
            self.wfile.write(raw)
        except Exception as e:
            self._json({"error": str(e)}, 500)

    def _lookups(self):
        con = connect()
        provs = [
            r[0]
            for r in con.execute(
                "SELECT DISTINCT province FROM charity_base ORDER BY 1"
            ).fetchall()
        ]
        desigs = [
            {"code": r[0], "desc": r[1]}
            for r in con.execute(
                "SELECT designation_code, designation_desc FROM charity_base GROUP BY 1,2 ORDER BY 1"
            ).fetchall()
        ]
        cats = [
            {"code": r[0], "desc": r[1]}
            for r in con.execute(
                "SELECT DISTINCT category_code, category_desc FROM charity_base "
                "WHERE category_desc IS NOT NULL ORDER BY 2"
            ).fetchall()
        ]
        con.close()
        self._json({"provinces": provs, "designations": desigs, "categories": cats})

    def _json(self, data, status=200):
        raw = json.dumps(data, default=default_ser).encode()
        self.send_response(status)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(raw)))
        self.end_headers()
        self.wfile.write(raw)

    def log_message(self, fmt, *args):
        if "/api/" not in str(args[0] if args else ""):
            super().log_message(fmt, *args)


def main():
    if not DB.exists():
        print(f"Database not found: {DB}")
        print("Run: python3 scripts/load_csv.py")
        sys.exit(1)

    srv = http.server.HTTPServer(("127.0.0.1", PORT), Handler)
    print(f"T3010 Data Explorer -> http://localhost:{PORT}")
    print("Press Ctrl+C to stop.\n")

    def open_browser():
        try:
            webbrowser.open(f"http://localhost:{PORT}")
        except Exception:
            pass

    threading.Timer(0.5, open_browser).start()

    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        print("\nStopped.")
        srv.shutdown()


if __name__ == "__main__":
    main()
